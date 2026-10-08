from pathlib import Path
import csv, hashlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parent
PAPER=BASE.parents[1]
DTA=PAPER/"data/UKDA-8971-stata/stata/stata13/dcms_csbs_combined_2016-2022_banded_cost_data_only_v3_public.dta"
BPE2019=PAPER/"data/external_candidates/BPE_2019_detailed_tables.xls"
BPE2020=PAPER/"data/external_candidates/BPE_2020_detailed_tables.xlsx"
BPE2021=PAPER/"data/external_candidates/BPE_2021_detailed_tables.xlsx"
G1E=PAPER/"4_Methodology_Design/G1_8971_Audit/G1E_INTEGRATED_ARCHITECTURE"

ATTACK=[f"type{i}" for i in list(range(1,10))+[13,15,16]]
P_ITEMS=["rules2","rules3","rules4","rules5","rules7","rules13","rules14","rules15","manage3","ident4"]
MAT=[f"outcome{i}" for i in [1,2,3,4,5,6,8]]
SECTOR_GROUPS={1:["L","N"],2:["F"],3:["P"],4:["R","S"],5:["K"],6:["I"],7:["Q"],8:["J"],9:["M"],10:["G"],11:["H"],12:["BDE","C"]}
STARTS={"A Agriculture":"A","B, D and E":"BDE","C Manufacturing":"C","F Construction":"F","G Wholesale":"G","H Transportation":"H","I Accommodation":"I","J Information":"J","K Financial":"K","L Real":"L","M Professional":"M","N Administrative":"N","P Education":"P","Q Human":"Q","R Arts":"R","S Other":"S"}

def sha(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""): h.update(b)
    return h.hexdigest()

def rake(g,size_prop,sector_prop):
    size=g.sizeb.astype(int).to_numpy(); sector=g.sector_comb2.astype(int).to_numpy()
    w=np.ones(len(g),float)
    for it in range(1000):
        old=w.copy()
        for k,p in size_prop.items():
            m=size==k; w[m]*=(p*w.sum()/w[m].sum())
        for k,p in sector_prop.items():
            m=sector==k; w[m]*=(p*w.sum()/w[m].sum())
        w*=len(g)/w.sum()
        if np.max(np.abs(w-old))<1e-12: return w,it+1
    raise RuntimeError("RIM failed to converge")

def parse_bpe_xlsx(path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb["Table 5"]
    sections={}; rows={}; cur=None
    for row in ws.iter_rows(values_only=True):
        first=row[0]
        if isinstance(first,str):
            for prefix,code in STARTS.items():
                if first.startswith(prefix): cur=code; rows.setdefault(code,{}); break
        if cur and isinstance(first,str) and first.strip() in ["1","2-4","5-9","10-19","20-49","50-99","100-199","200-249","250-499","500 or more"]:
            rows[cur][first.strip()]=float(row[2])
        if cur and first=="All employers": sections[cur]=float(row[2])
    return target_props(rows,sections)

def parse_bpe2019_xls(path):
    d=pd.read_excel(path,sheet_name="Table 5",header=None,engine="xlrd")
    sections={}; rows={}; cur=None
    for _,row in d.iterrows():
        first=row.iloc[0]
        if isinstance(first,str):
            for prefix,code in STARTS.items():
                if first.startswith(prefix): cur=code; rows.setdefault(code,{}); break
        if cur and isinstance(first,str) and first.strip() in ["1","2-4","5-9","10-19","20-49","50-99","100-199","200-249","250-499","500 or more"]:
            rows[cur][first.strip()]=float(row.iloc[2])
        if cur and first=="All employers": sections[cur]=float(row.iloc[2])
    return target_props(rows,sections)

def target_props(rows,sections):
    codes=[c for cs in SECTOR_GROUPS.values() for c in cs]
    s={1:sum(rows[c]["1"]+rows[c]["2-4"]+rows[c]["5-9"] for c in codes),
       2:sum(rows[c]["10-19"]+rows[c]["20-49"] for c in codes),
       3:sum(rows[c]["50-99"]+rows[c]["100-199"]+rows[c]["200-249"] for c in codes),
       4:sum(rows[c]["250-499"]+rows[c]["500 or more"] for c in codes)}
    c={k:sum(sections[z] for z in zs) for k,zs in SECTOR_GROUPS.items()}
    return {k:v/sum(s.values()) for k,v in s.items()},{k:v/sum(c.values()) for k,v in c.items()}

df=pd.read_stata(DTA,convert_categoricals=False)

# 2019 common calibration: exact BPE 2018 four-size margins from official BPE 2019 time series.
size2018={1:146200+734100+257000,2:137400+72200,3:23000+9800+2000,4:3800+3700}
sp19={k:v/sum(size2018.values()) for k,v in size2018.items()}
g19=df[(df.year==4)&(df.questtype==1)&(df.sector_comb2.isin(range(1,13)))&(df.weight>0)].copy()
ow19=g19.weight.to_numpy(float); sec19=g19.sector_comb2.astype(int).to_numpy()
cp19={k:ow19[sec19==k].sum()/ow19.sum() for k in range(1,13)}
w19,it19=rake(g19,sp19,cp19)

# 2020 common calibration: exact BPE 2019 employer margins.
sp20,cp20=parse_bpe2019_xls(BPE2019)
g20=df[(df.year==5)&(df.questtype==1)&(df.sector_comb2.isin(range(1,13)))].copy()
w20,it20=rake(g20,sp20,cp20)

# 2022 validation weights use all 13 sectors. Primary common population excludes agriculture,
# then rakes to exact BPE 2021 employer margins for sectors 1-12.
sp22,cp22=parse_bpe_xlsx(BPE2021)
g22=df[(df.year==7)&(df.questtype==1)&(df.sector_comb2.isin(range(1,13)))].copy()
w22,it22=rake(g22,sp22,cp22)

# Build crosswalk.
parts=[]
for year,g,w,source in [
    (2019,g19,w19,"retrospective_common_calibration_BPE2018_size_verified2019_sector"),
    (2020,g20,w20,"official_method_reconstruction_BPE2019"),
    (2022,g22,w22,"common12_reconstruction_BPE2021"),
]:
    if year==2019:
        legacy=g.weight.to_numpy(float)
        legacy=legacy/legacy.mean()
    else:
        legacy=w.copy()
    q=pd.DataFrame({
        "uniqser":g.uniqser.astype(str),
        "target_year":year,
        "analysis_weight_common12":w,
        "sensitivity_weight_annual_method":legacy,
        "weight_source":source,
    })
    parts.append(q)
cross=pd.concat(parts,ignore_index=True)
assert not cross.duplicated(["uniqser","target_year"]).any()
cross.to_csv(BASE/"G1F_COMMON12_WEIGHT_CROSSWALK.csv",index=False)

# Diagnostics using attack occurrence only as external validation, never for weight fitting.
rows=[]
for year,g,w,published in [(2019,g19,w19,0.32),(2020,g20,w20,0.46),(2022,g22,w22,None)]:
    attack=g[ATTACK].eq(1).any(axis=1).to_numpy(float)
    rows.append({"year":year,"n":len(g),"weight_sum":w.sum(),"weight_min":w.min(),"weight_max":w.max(),
                 "kish_effective_n":w.sum()**2/(w@w),"attack_rate_weighted":np.average(attack,weights=w),
                 "published_all_sector_attack_rate":published})
pd.DataFrame(rows).to_csv(BASE/"G1F_COMMON12_WEIGHT_DIAGNOSTICS.csv",index=False)

# 2022 independent all-sector validation including agriculture.
# Parse BPE2021 again including agriculture for exact size and 13-sector margins.
wb=load_workbook(BPE2021,read_only=True,data_only=True); ws=wb["Table 5"]
sections={}; rows_b={}; cur=None
for row in ws.iter_rows(values_only=True):
    first=row[0]
    if isinstance(first,str):
        for prefix,code in STARTS.items():
            if first.startswith(prefix): cur=code; rows_b.setdefault(code,{}); break
    if cur and isinstance(first,str) and first.strip() in ["1","2-4","5-9","10-19","20-49","50-99","100-199","200-249","250-499","500 or more"]:
        rows_b[cur][first.strip()]=float(row[2])
    if cur and first=="All employers": sections[cur]=float(row[2])
groups13={**SECTOR_GROUPS,13:["A"]}; codes13=[c for cs in groups13.values() for c in cs]
s13={1:sum(rows_b[c]["1"]+rows_b[c]["2-4"]+rows_b[c]["5-9"] for c in codes13),
     2:sum(rows_b[c]["10-19"]+rows_b[c]["20-49"] for c in codes13),
     3:sum(rows_b[c]["50-99"]+rows_b[c]["100-199"]+rows_b[c]["200-249"] for c in codes13),
     4:sum(rows_b[c]["250-499"]+rows_b[c]["500 or more"] for c in codes13)}
c13={k:sum(sections[z] for z in zs) for k,zs in groups13.items()}
sp13={k:v/sum(s13.values()) for k,v in s13.items()}; cp13={k:v/sum(c13.values()) for k,v in c13.items()}
g22all=df[(df.year==7)&(df.questtype==1)&(df.sector_comb2.isin(range(1,14)))].copy()
w22all,_=rake(g22all,sp13,cp13)
attack22=g22all[ATTACK].eq(1).any(axis=1).to_numpy(float)
with open(BASE/"G1F_2022_EXTERNAL_WEIGHT_VALIDATION.csv","w",newline="",encoding="utf-8") as f:
    z=csv.writer(f); z.writerow(["metric","reconstructed","published"])
    z.writerow(["attack_rate",np.average(attack22,weights=w22all),0.39])
    z.writerow(["kish_effective_n",w22all.sum()**2/(w22all@w22all),817])

# Source hashes.
with open(BASE/"G1F_WEIGHT_SOURCE_PROVENANCE.csv","w",newline="",encoding="utf-8") as f:
    z=csv.writer(f); z.writerow(["source","sha256"])
    for p in [DTA,BPE2019,BPE2020,BPE2021]:
        z.writerow([str(p.relative_to(PAPER)),sha(p)])

print(cross.groupby("target_year").size().to_string())
print(pd.read_csv(BASE/"G1F_COMMON12_WEIGHT_DIAGNOSTICS.csv").to_string(index=False))
print(pd.read_csv(BASE/"G1F_2022_EXTERNAL_WEIGHT_VALIDATION.csv").to_string(index=False))
